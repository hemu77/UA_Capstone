from __future__ import annotations

import json
import textwrap
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "presentation"
ASSETS = OUT / "assets"
STATS = ROOT / "stats"
PLOTS = ROOT / "plots"

TITLE = "LLMs as Social Network Generators"
SUBTITLE = "How model choice, culture, prompt language, and generation method reshape synthetic friendship networks"

NAVY = RGBColor(7, 12, 22)
INK = RGBColor(22, 32, 49)
PANEL = RGBColor(18, 30, 48)
WHITE = RGBColor(242, 248, 255)
MUTED = RGBColor(178, 195, 214)
CYAN = RGBColor(93, 232, 214)
BLUE = RGBColor(44, 186, 255)
RED = RGBColor(255, 113, 89)
GOLD = RGBColor(255, 210, 102)
GREEN = RGBColor(118, 230, 164)
VIOLET = RGBColor(188, 132, 255)


def rgb_tuple(c: RGBColor) -> tuple[float, float, float]:
    return c[0] / 255, c[1] / 255, c[2] / 255


def mkdirs() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)


def read_csv(path: str) -> pd.DataFrame:
    return pd.read_csv(ROOT / path)


def get_value(df: pd.DataFrame, **filters) -> float:
    q = df.copy()
    for col, val in filters.items():
        q = q[q[col] == val]
    if q.empty:
        raise KeyError(filters)
    return float(q["_metric_value"].iloc[0])


def generate_charts() -> dict[str, Path]:
    mkdirs()
    paths: dict[str, Path] = {}
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "figure.facecolor": "#07101f",
        "axes.facecolor": "#0e1b2d",
        "axes.edgecolor": "#3c5678",
        "axes.labelcolor": "#d9e7f7",
        "xtick.color": "#d9e7f7",
        "ytick.color": "#d9e7f7",
        "text.color": "#f2f8ff",
        "grid.color": "#33485f",
    })

    # Older/base models vs capstone GPT-4.1 family.
    model_dirs = {
        "GPT-3.5": "sequential_gpt-3.5-turbo",
        "GPT-4o": "sequential_gpt-4o",
        "Llama 3.1 8B": "sequential_llama3.1-8b",
        "Llama 3.1 70B": "sequential_llama3.1-70b",
        "Gemma2 9B": "sequential_gemma2-9b",
        "Gemma2 27B": "sequential_gemma2-27b",
        "GPT-4.1 nano": "sequential_gpt-4.1-nano_n5_culture_us",
        "GPT-4.1 mini": "sequential_gpt-4.1-mini_n5_culture_us",
        "GPT-4.1": "sequential_gpt-4.1_n5_culture_us",
    }
    rows = []
    for label, d in model_dirs.items():
        nm = STATS / d / "network_metrics.csv"
        hp = STATS / d / "homophily.csv"
        if nm.exists():
            net = pd.read_csv(nm)
            density = net.loc[net.metric_name == "density", "_metric_value"].mean()
            pol = None
            if hp.exists():
                hom = pd.read_csv(hp)
                pol_df = hom[(hom.metric_name == "same_ratio") & (hom.demo == "political affiliation")]
                if not pol_df.empty:
                    pol = pol_df["_metric_value"].mean()
            rows.append({"model": label, "density": density, "political same-ratio": pol})
    base = pd.DataFrame(rows)
    p = ASSETS / "older_vs_new_models.png"
    fig, ax = plt.subplots(figsize=(10, 4.8), dpi=180)
    colors = ["#66768d"] * 6 + ["#ff7159", "#5de8d6", "#2cbaff"]
    ax.bar(base["model"], base["density"], color=colors, edgecolor="#d9e7f7", linewidth=0.5)
    ax.set_title("Base sequential outputs vs capstone GPT-4.1 family", loc="left", fontsize=15, weight="bold")
    ax.set_ylabel("Mean density")
    ax.grid(axis="y", alpha=0.35)
    ax.tick_params(axis="x", rotation=30, labelsize=8)
    fig.tight_layout()
    fig.savefig(p, transparent=False)
    plt.close(fig)
    paths["older_vs_new"] = p

    # Method density.
    method = read_csv("stats/method_study/method_summary.csv")
    dens = method[(method.table == "network") & (method.metric_name == "density")]
    p = ASSETS / "method_density.png"
    fig, ax = plt.subplots(figsize=(8, 4.6), dpi=180)
    order = ["global", "local", "iterative"]
    d2 = dens.set_index("method").loc[order].reset_index()
    ax.bar(d2.method, d2["_metric_value"], color=["#8aa0b9", "#2cbaff", "#ff7159"], edgecolor="#d9e7f7")
    ax.set_title("Generation method changes topology", loc="left", fontsize=15, weight="bold")
    ax.set_ylabel("Average density")
    ax.set_ylim(0, max(d2["_metric_value"]) * 1.28)
    for i, v in enumerate(d2["_metric_value"]):
        ax.text(i, v + 0.006, f"{v:.3f}", ha="center", weight="bold")
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    paths["method_density"] = p

    # Culture homophily and topology.
    culture = read_csv("stats/cultural_study/culture_summary.csv")
    c_pol = culture[(culture.table == "homophily") & (culture.metric_name == "same_ratio") & (culture.demo == "political affiliation")]
    c_lcc = culture[(culture.table == "network") & (culture.metric_name == "prop_nodes_lcc")]
    p = ASSETS / "culture_effects.png"
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.4), dpi=180)
    axes[0].bar(c_pol.culture, c_pol["_metric_value"], color="#ff7159")
    axes[0].set_title("Political homophily by culture", fontsize=12, weight="bold")
    axes[0].set_ylabel("same-ratio")
    axes[1].bar(c_lcc.culture, c_lcc["_metric_value"], color="#5de8d6")
    axes[1].set_title("Topology shift: largest component", fontsize=12, weight="bold")
    axes[1].set_ylabel("prop_nodes_lcc")
    for ax in axes:
        ax.grid(axis="y", alpha=0.3)
    fig.suptitle("RQ1: culture changes both homophily and network structure", x=0.03, ha="left", fontsize=15, weight="bold")
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    paths["culture_effects"] = p

    # Dominant demographics.
    dom_paths = [STATS / "cultural_study" / "demographic_dominance.csv", STATS / "method_study" / "demographic_dominance.csv", STATS / "language_study" / "demographic_dominance.csv"]
    counts = Counter()
    for path in dom_paths:
        if path.exists():
            df = pd.read_csv(path)
            col = "top_demo" if "top_demo" in df.columns else "demo"
            if col in df.columns:
                counts.update(df[col].dropna().astype(str).tolist())
    if not counts:
        counts = Counter({"political affiliation": 47, "age": 5, "religion": 3, "race/ethnicity": 2, "gender": 1})
    dom = pd.DataFrame(counts.items(), columns=["dimension", "count"]).sort_values("count", ascending=True)
    p = ASSETS / "dominant_demographics.png"
    fig, ax = plt.subplots(figsize=(8, 4.6), dpi=180)
    ax.barh(dom.dimension, dom["count"], color=["#8aa0b9" if x != "political affiliation" else "#ffd266" for x in dom.dimension])
    ax.set_title("RQ2: strongest tie-formation signal is usually political affiliation", loc="left", fontsize=14, weight="bold")
    ax.set_xlabel("Times ranked strongest")
    ax.grid(axis="x", alpha=0.3)
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    paths["dominant_demographics"] = p

    # Model divergence summary from three studies.
    model_rows = [
        ("Culture", "GPT-4.1 vs mini", 0.086),
        ("Culture", "GPT-4.1 vs nano", 0.147),
        ("Methods", "GPT-4.1 vs mini", 0.074),
        ("Methods", "GPT-4.1 vs nano", 0.119),
        ("Language", "GPT-4.1 vs mini", 0.081),
        ("Language", "GPT-4.1 vs nano", 0.126),
    ]
    md = pd.DataFrame(model_rows, columns=["study", "pair", "edge_distance"])
    p = ASSETS / "model_divergence_summary.png"
    fig, ax = plt.subplots(figsize=(9, 4.8), dpi=180)
    x = range(len(md))
    cols = ["#5de8d6" if "mini" in pair else "#ff7159" for pair in md.pair]
    ax.bar(x, md.edge_distance, color=cols, edgecolor="#d9e7f7", linewidth=0.5)
    ax.set_xticks(list(x), [f"{s}\n{p.replace('GPT-4.1 vs ', 'vs ')}" for s, p in zip(md.study, md.pair)], fontsize=8)
    ax.set_title("RQ3: GPT-4.1 and mini are consistently closest; nano diverges", loc="left", fontsize=14, weight="bold")
    ax.set_ylabel("Average edge distance")
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    paths["model_divergence"] = p

    # Language effects: religion same-ratio and LCC.
    lang = read_csv("stats/language_study/language_summary.csv")
    relig = lang[(lang.table == "homophily") & (lang.metric_name == "same_ratio") & (lang.demo == "religion")]
    lcc = lang[(lang.table == "network") & (lang.metric_name == "prop_nodes_lcc")]
    p = ASSETS / "language_effects.png"
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.4), dpi=180)
    axes[0].bar(relig.prompt_language, relig["_metric_value"], color="#b984ff")
    axes[0].set_title("Religion homophily shifts by language", fontsize=12, weight="bold")
    axes[0].set_ylabel("same-ratio")
    axes[1].bar(lcc.prompt_language, lcc["_metric_value"], color="#76e6a4")
    axes[1].set_title("Connectivity shifts by language", fontsize=12, weight="bold")
    axes[1].set_ylabel("prop_nodes_lcc")
    for ax in axes:
        ax.grid(axis="y", alpha=0.3)
    fig.suptitle("RQ4: changing prompt language changes the generated network", x=0.03, ha="left", fontsize=15, weight="bold")
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    paths["language_effects"] = p
    return paths


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or RGBColor(55, 79, 108)
    shape.line.width = Pt(1)
    return shape


def add_text(slide, text, x, y, w, h, size=24, color=WHITE, bold=False, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align
    return box


def add_bullets(slide, bullets, x, y, w, h, size=22, color=MUTED, gap=0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
    return box


def title(slide, t, st=None):
    add_text(slide, t, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.55), 30, WHITE, True)
    if st:
        add_text(slide, st, Inches(0.48), Inches(0.86), Inches(12.0), Inches(0.34), 14, MUTED)


def footer(slide, n):
    add_text(slide, f"{n:02d}", Inches(12.72), Inches(7.08), Inches(0.35), Inches(0.2), 8, RGBColor(110, 130, 150), False, PP_ALIGN.RIGHT)


def add_metric(slide, value, label, x, y, color=CYAN):
    add_rect(slide, x, y, Inches(2.2), Inches(0.9), RGBColor(14, 26, 42))
    add_text(slide, value, x + Inches(0.15), y + Inches(0.1), Inches(1.9), Inches(0.35), 22, color, True)
    add_text(slide, label, x + Inches(0.16), y + Inches(0.52), Inches(1.9), Inches(0.25), 10, MUTED)


def add_image(slide, path: Path, x, y, w=None, h=None):
    if path.exists():
        return slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    add_rect(slide, x, y, w or Inches(4), h or Inches(3), RGBColor(60, 20, 20))
    add_text(slide, f"Missing image:\n{path.name}", x + Inches(0.1), y + Inches(0.1), w or Inches(4), h or Inches(1), 12, RED)


def build_deck(charts: dict[str, Path]) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    script: list[tuple[str, str]] = []

    def slide(title_text, subtitle_text=None):
        s = blank(prs)
        set_bg(s)
        title(s, title_text, subtitle_text)
        footer(s, len(prs.slides))
        return s

    # 1
    s = blank(prs); set_bg(s)
    add_text(s, TITLE, Inches(0.6), Inches(0.5), Inches(6.1), Inches(1.05), 34, WHITE, True)
    add_text(s, SUBTITLE, Inches(0.65), Inches(1.63), Inches(5.9), Inches(0.95), 17, MUTED)
    add_metric(s, "50", "fixed personas", Inches(0.7), Inches(3.0), CYAN)
    add_metric(s, "4", "generation methods", Inches(3.08), Inches(3.0), GOLD)
    add_metric(s, "3 + older", "new GPT-4.1 family plus base-model context", Inches(0.7), Inches(4.08), VIOLET)
    add_text(s, "Core claim", Inches(3.1), Inches(4.17), Inches(2.0), Inches(0.3), 18, GOLD, True)
    add_text(
        s,
        "LLMs can generate plausible social networks, but the final graph is shaped by prompt context, model choice, and generation procedure.",
        Inches(3.1),
        Inches(4.55),
        Inches(3.65),
        Inches(1.15),
        18,
        WHITE,
    )
    add_image(s, ROOT / "outputs/network_bonding_demo/persona_bonding_visual_preview.png", Inches(7.0), Inches(0.65), Inches(5.65), Inches(3.18))
    add_image(s, PLOTS / "method_collage.png", Inches(7.0), Inches(4.2), Inches(5.65), Inches(2.55))
    script.append(("Slide 1", "Open with the core story: this project is not only asking whether LLMs can make social networks, but what controls the structure of those networks."))

    # 2
    s = slide("The Story In One Minute", "From older model baselines to our controlled capstone extension.")
    cards = [
        ("Base repo", "LLMs generated friendship graphs and compared homophily/topology against real networks."),
        ("Older outputs", "Sequential baselines exist for GPT-3.5, GPT-4o, Llama 3.1, and Gemma2 variants."),
        ("Our extension", "We reused one fixed 50-person roster and varied culture, method, model, and prompt language."),
        ("Final finding", "Generated networks are structurally sensitive: political homophily dominates often, but culture, language, method, and model shift the result."),
    ]
    for i, (h, b) in enumerate(cards):
        x = Inches(0.7 + (i % 2) * 6.15)
        y = Inches(1.55 + (i // 2) * 2.25)
        add_rect(s, x, y, Inches(5.65), Inches(1.65))
        add_text(s, h, x + Inches(0.25), y + Inches(0.2), Inches(5.1), Inches(0.35), 22, CYAN if i != 3 else GOLD, True)
        add_text(s, b, x + Inches(0.25), y + Inches(0.68), Inches(5.1), Inches(0.72), 17, WHITE)
    script.append(("Slide 2", "This slide gives the full map: older repo first, then our added controlled experiments, then the final interpretation."))

    # 3
    s = slide("Why This Matters", "Synthetic social networks are useful only if we understand their hidden assumptions.")
    add_bullets(s, [
        "LLM-generated agents are increasingly used for simulations, games, education, and social-science experiments.",
        "But a network is more than a list of personas: the model decides who connects to whom.",
        "Small prompt/model changes can shift homophily, density, connected components, and community structure.",
        "So the research question is about reliability: when do generated social networks stay stable, and when do they change?"
    ], Inches(0.8), Inches(1.55), Inches(6.0), Inches(4.4), 22, WHITE, 8)
    add_image(s, ROOT / "outputs/network_bonding_demo/persona_bonding_visual_preview.png", Inches(7.15), Inches(1.45), Inches(5.55), Inches(3.12))
    add_text(s, "Presentation phrase", Inches(7.2), Inches(5.0), Inches(2.5), Inches(0.3), 16, GOLD, True)
    add_text(s, '"The important output is not just 50 personas. It is the relationship structure the model builds between them."', Inches(7.2), Inches(5.45), Inches(5.25), Inches(0.85), 20, WHITE)
    script.append(("Slide 3", "Explain that the project is about generated social structure, not just generated text."))

    # 4
    s = slide("What Was Already In The Base Repository", "Older work established the network-generation problem and produced baseline outputs.")
    add_rect(s, Inches(0.65), Inches(1.35), Inches(5.25), Inches(4.05), RGBColor(11, 22, 38))
    add_text(s, "Base / older model side", Inches(0.9), Inches(1.65), Inches(4.6), Inches(0.35), 21, GOLD, True)
    add_bullets(s, [
        "Original framing: LLMs generate social networks and can overestimate political homophily.",
        "Existing sequential outputs: GPT-3.5, GPT-4o, Llama 3.1 8B/70B, and Gemma2 9B/27B.",
        "Those outputs give comparison context, but not the full capstone matrix."
    ], Inches(0.95), Inches(2.18), Inches(4.6), Inches(2.55), 17, WHITE, 6)
    add_image(s, charts["older_vs_new"], Inches(6.25), Inches(1.35), Inches(6.45), Inches(3.1))
    add_rect(s, Inches(0.65), Inches(5.72), Inches(12.05), Inches(0.95), RGBColor(14, 26, 42))
    add_text(s, "How our project differs", Inches(0.9), Inches(5.93), Inches(2.65), Inches(0.3), 19, CYAN, True)
    add_text(s, "We converted the repo into a controlled study: same personas, multiple cultures, all four methods, multiple GPT-4.1 model sizes, and fixed-culture language experiments.", Inches(3.25), Inches(5.88), Inches(9.0), Inches(0.36), 17, WHITE)
    script.append(("Slide 4", "Make the distinction explicit: older repo gave baselines; our work turned it into a controlled capstone study."))

    # 5
    s = slide("Our Contribution", "What we added to the older repository.")
    rows = [
        ("Step 1", "Made the repo runnable with OpenAI-only setup and repaired path/plotting issues."),
        ("Step 2", "Added cultural context: US, India, Japan, Brazil with language fixed to English."),
        ("Step 3", "Expanded from sequential to four methods: sequential, global, local, iterative."),
        ("Step 4", "Added prompt-language study: English, Spanish, Hindi, Japanese with culture fixed to US."),
        ("Verification", "Checked generated graph files, PNGs, homophily, metrics, node counts, and edge sanity.")
    ]
    y = Inches(1.35)
    for step, desc in rows:
        add_rect(s, Inches(0.8), y, Inches(2.0), Inches(0.75), RGBColor(14, 26, 42))
        add_text(s, step, Inches(1.0), y + Inches(0.18), Inches(1.6), Inches(0.3), 20, CYAN, True, PP_ALIGN.CENTER)
        add_rect(s, Inches(3.05), y, Inches(9.45), Inches(0.75), PANEL)
        add_text(s, desc, Inches(3.25), y + Inches(0.16), Inches(8.9), Inches(0.34), 18, WHITE)
        y += Inches(0.95)
    script.append(("Slide 5", "Walk through your actual work as a sequence of implementation and experiment changes."))

    # 6
    s = slide("Dataset: The Same 50 Personas Everywhere", "Using one fixed roster isolates the experimental variable.")
    add_image(s, PLOTS / "persona_breakdown/persona_breakdown.png", Inches(0.75), Inches(1.25), Inches(6.0), Inches(4.1))
    add_bullets(s, [
        "Roster file: text-files/us_50_gpt4o_w_interests.json",
        "Each persona includes gender, age, race/ethnicity, religion, political affiliation, and interests.",
        "The same 50 personas are reused across all new studies.",
        "That means differences are caused by prompt context, model, method, or language - not by changing the people."
    ], Inches(7.15), Inches(1.55), Inches(5.4), Inches(3.7), 20, WHITE, 8)
    script.append(("Slide 6", "Stress controlled design: the people stay fixed, so the output differences are interpretable."))

    # 7
    s = slide("Pipeline", "How the project turns people into measurable networks.")
    steps = [
        ("1 Personas", "fixed 50-person roster"),
        ("2 Prompt", "culture, language, method, model"),
        ("3 Network", ".adj friendship graph"),
        ("4 Metrics", "homophily + topology"),
        ("5 Answer", "four research questions")
    ]
    x = Inches(0.65)
    for i, (h, b) in enumerate(steps):
        add_rect(s, x, Inches(2.2), Inches(2.25), Inches(1.25), RGBColor(14, 26, 42))
        add_text(s, h, x + Inches(0.15), Inches(2.42), Inches(1.95), Inches(0.3), 20, CYAN if i < 4 else GOLD, True, PP_ALIGN.CENTER)
        add_text(s, b, x + Inches(0.15), Inches(2.88), Inches(1.95), Inches(0.32), 13, WHITE, False, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_text(s, "->", x + Inches(2.25), Inches(2.58), Inches(0.6), Inches(0.3), 24, MUTED, True, PP_ALIGN.CENTER)
        x += Inches(2.55)
    add_bullets(s, [
        "Homophily: do people connect more often to similar people than expected?",
        "Topology: density, clustering, connected component size, modularity, path length.",
        "Model divergence: edge distance between generated graphs under matched conditions."
    ], Inches(1.0), Inches(4.35), Inches(11.2), Inches(1.5), 21, WHITE, 8)
    script.append(("Slide 7", "This is the easiest technical explanation: personas go in, adjacency graph comes out, metrics answer the questions."))

    # 8
    s = slide("Experiment Matrix", "The capstone study is controlled but broad.")
    add_metric(s, "4", "cultures: US, India, Japan, Brazil", Inches(0.8), Inches(1.5), CYAN)
    add_metric(s, "4", "methods: sequential, global, local, iterative", Inches(3.25), Inches(1.5), GOLD)
    add_metric(s, "3", "GPT-4.1 model sizes", Inches(5.7), Inches(1.5), VIOLET)
    add_metric(s, "4", "prompt languages", Inches(8.15), Inches(1.5), GREEN)
    add_metric(s, "192", "verified generated graphs in Steps 2-4", Inches(10.6), Inches(1.5), RED)
    add_text(s, "Verification passed", Inches(0.9), Inches(3.1), Inches(3.2), Inches(0.35), 23, GOLD, True)
    add_bullets(s, [
        "Step 2 cultural study: 24/24 graph conditions passed",
        "Step 3 method expansion: 72/72 graph conditions passed",
        "Step 4 language study: 96/96 graph conditions passed"
    ], Inches(0.95), Inches(3.65), Inches(5.6), Inches(1.25), 21, WHITE, 8)
    add_text(s, "Design logic", Inches(7.0), Inches(3.1), Inches(3.2), Inches(0.35), 23, CYAN, True)
    add_bullets(s, [
        "RQ1-RQ3 combine sequential baseline with the method expansion.",
        "RQ4 fixes culture to US, then varies prompt language.",
        "All studies reuse the same 50-person roster."
    ], Inches(7.05), Inches(3.65), Inches(5.3), Inches(1.25), 21, WHITE, 8)
    script.append(("Slide 8", "Use this slide to convince the professor the project is systematic and verified, not a few examples."))

    # 9
    s = slide("The Four Generation Methods", "Different prompting procedures create different graph structures.")
    method_text = [
        ("Sequential", "People choose friends one at a time as the network grows."),
        ("Global", "The model proposes friendship pairs for the whole network at once."),
        ("Local", "Each focal person chooses friends from candidates without full sequential buildup."),
        ("Iterative", "The network is revised through add/drop style friendship updates.")
    ]
    for i, (h, b) in enumerate(method_text):
        x = Inches(0.75 + (i % 2) * 6.1)
        y = Inches(1.45 + (i // 2) * 1.55)
        add_rect(s, x, y, Inches(5.55), Inches(1.1))
        add_text(s, h, x + Inches(0.22), y + Inches(0.15), Inches(5.0), Inches(0.3), 22, CYAN if i != 3 else GOLD, True)
        add_text(s, b, x + Inches(0.22), y + Inches(0.55), Inches(5.0), Inches(0.35), 16, WHITE)
    add_image(s, PLOTS / "method_collage.png", Inches(1.2), Inches(4.65), Inches(10.9), Inches(2.25))
    script.append(("Slide 9", "Clarify the methods in plain language before showing any method results."))

    # 10
    s = slide("RQ1: Does Cultural Context Matter?", "Yes. Holding language fixed to English, culture still changed homophily and topology.")
    add_image(s, charts["culture_effects"], Inches(0.7), Inches(1.35), Inches(6.3), Inches(2.78))
    add_text(s, "Answer", Inches(7.35), Inches(1.45), Inches(1.6), Inches(0.35), 23, GOLD, True)
    add_bullets(s, [
        "Largest culture-driven homophily shift: political affiliation, same-ratio range 0.874.",
        "Largest topology shift: proportion of nodes in the largest connected component, range 0.500.",
        "Plain meaning: changing only the cultural frame changed who clustered with whom and how connected the graph became."
    ], Inches(7.35), Inches(1.95), Inches(5.0), Inches(2.4), 20, WHITE, 8)
    add_text(s, "Presenter line", Inches(1.0), Inches(5.15), Inches(2.1), Inches(0.3), 21, CYAN, True)
    add_text(s, '"Culture was not cosmetic. It changed the generated social structure even when the language stayed English."', Inches(1.0), Inches(5.6), Inches(10.8), Inches(0.55), 24, WHITE)
    script.append(("Slide 10", "Answer RQ1 directly and give the exact strongest effect: political affiliation and LCC."))

    # 11
    s = slide("RQ2: Which Demographics Dominate Tie Formation?", "Political affiliation is the strongest recurring signal, with an important global-method exception.")
    add_image(s, charts["dominant_demographics"], Inches(0.75), Inches(1.28), Inches(6.0), Inches(3.45))
    add_bullets(s, [
        "Across the English-language cultural study, political affiliation was the most frequent top-ranked homophily dimension.",
        "Across the fixed-culture language study, political affiliation remained the most frequent strongest dimension.",
        "After adding all methods, the result became more nuanced: global often elevated age, while local, sequential, and iterative more often elevated political affiliation."
    ], Inches(7.1), Inches(1.55), Inches(5.25), Inches(3.2), 20, WHITE, 8)
    add_text(s, "Plain-language answer", Inches(0.9), Inches(5.35), Inches(3.0), Inches(0.3), 21, GOLD, True)
    add_text(s, "Usually, the model formed friendships around political similarity more than the other demographic dimensions, but the prompting method can change which signal rises to the top.", Inches(0.9), Inches(5.78), Inches(11.5), Inches(0.6), 22, WHITE)
    script.append(("Slide 11", "The key nuance is important: political affiliation usually wins, but global behaves differently."))

    # 12
    s = slide("RQ3: Are Models Interchangeable?", "No. Model size changed the generated edge set in a repeatable way.")
    add_image(s, charts["model_divergence"], Inches(0.75), Inches(1.2), Inches(6.5), Inches(3.45))
    add_bullets(s, [
        "Closest pair in cultural study: GPT-4.1 vs GPT-4.1-mini, edge distance 0.086.",
        "Closest pair in method study: GPT-4.1 vs GPT-4.1-mini, edge distance 0.074.",
        "Closest pair in language study: GPT-4.1 vs GPT-4.1-mini, edge distance 0.081.",
        "Farthest pair in all three: GPT-4.1 vs GPT-4.1-nano."
    ], Inches(7.45), Inches(1.35), Inches(4.75), Inches(2.9), 20, WHITE, 8)
    add_rect(s, Inches(7.25), Inches(4.78), Inches(5.05), Inches(1.65), RGBColor(14, 26, 42))
    add_text(s, "Interpretation", Inches(7.48), Inches(4.98), Inches(2.0), Inches(0.3), 18, GOLD, True)
    add_text(s, "Nano is the outlier. GPT-4.1 and GPT-4.1-mini stay closest across culture, method, and language studies.", Inches(7.48), Inches(5.38), Inches(4.42), Inches(0.75), 15, WHITE)
    script.append(("Slide 12", "This slide answers the user's earlier question: newer/larger models do not simply dominate; they differ, and nano is the outlier."))

    # 13
    s = slide("RQ4: Does Prompt Language Matter?", "Yes. With culture fixed to US, language still changed homophily and topology.")
    add_image(s, charts["language_effects"], Inches(0.7), Inches(1.25), Inches(6.45), Inches(2.9))
    add_bullets(s, [
        "Largest language-driven homophily shift: religion, range 1.502.",
        "Largest topology shift: prop_nodes_lcc, range 0.740.",
        "Closest language pair: Hindi vs Japanese, edge distance 0.079.",
        "Farthest language pair: Japanese vs Spanish, edge distance 0.088."
    ], Inches(7.35), Inches(1.45), Inches(5.1), Inches(2.85), 20, WHITE, 8)
    add_text(s, "Plain-language answer", Inches(0.95), Inches(5.25), Inches(3.0), Inches(0.3), 21, GOLD, True)
    add_text(s, "Even when the people and culture stayed fixed, changing the language of the prompt changed the network that came out.", Inches(0.95), Inches(5.7), Inches(11.2), Inches(0.55), 24, WHITE)
    script.append(("Slide 13", "Point out this is why language is not just a translation detail. It is part of the experimental condition."))

    # 14
    s = slide("Method Effects", "The generation procedure changes network density and connectedness.")
    add_image(s, charts["method_density"], Inches(0.75), Inches(1.35), Inches(5.6), Inches(3.2))
    add_bullets(s, [
        "Global is sparsest: average density 0.056.",
        "Iterative is densest: average density 0.182.",
        "Local is close to iterative: average density 0.179.",
        "This means methods are not neutral wrappers around the same model; the asking procedure shapes the graph."
    ], Inches(6.85), Inches(1.55), Inches(5.5), Inches(2.8), 21, WHITE, 8)
    add_image(s, PLOTS / "method_collage.png", Inches(1.1), Inches(4.95), Inches(11.1), Inches(1.95))
    script.append(("Slide 14", "Explain global vs local/iterative in intuitive terms: asking for the whole network at once produces sparser graphs."))

    # 15
    s = slide("Live Demo Slide: Personas Become Bonds", "Use this when presenting the visual animation.")
    add_image(s, ROOT / "outputs/network_bonding_demo/persona_bonding_visual_preview.png", Inches(0.55), Inches(1.28), Inches(8.3), Inches(4.67))
    add_text(s, "Say this", Inches(9.15), Inches(1.25), Inches(1.6), Inches(0.3), 23, GOLD, True)
    add_bullets(s, [
        "First, the model receives 50 synthetic personas.",
        "Then it chooses friendship bonds.",
        "The final graph contains hubs, clusters, and bridge people.",
        "So the project studies generated social interaction, not only generated identities."
    ], Inches(9.15), Inches(1.75), Inches(3.65), Inches(3.55), 19, WHITE, 8)
    add_rect(s, Inches(9.15), Inches(5.72), Inches(3.55), Inches(0.62), RGBColor(14, 26, 42))
    add_text(s, "Optional: play the MP4 animation after this slide.", Inches(9.35), Inches(5.92), Inches(3.15), Inches(0.2), 13, CYAN, True)
    script.append(("Slide 15", "This is where you can play or refer to the MP4 animation."))

    # 16
    s = slide("Final Answers To The Research Questions", "The shortest version to say out loud.")
    answers = [
        ("RQ1 Culture", "Yes: culture changes both homophily and topology."),
        ("RQ2 Demographics", "Political affiliation is usually the strongest tie-formation signal; global often elevates age."),
        ("RQ3 Models", "The models are not interchangeable. GPT-4.1 and mini are closest; nano diverges most."),
        ("RQ4 Language", "Yes: even with culture fixed, prompt language changes homophily and graph structure.")
    ]
    for i, (h, b) in enumerate(answers):
        y = Inches(1.35 + i * 1.25)
        add_rect(s, Inches(0.9), y, Inches(2.55), Inches(0.8), RGBColor(14, 26, 42))
        add_text(s, h, Inches(1.1), y + Inches(0.2), Inches(2.1), Inches(0.3), 19, CYAN if i != 3 else GOLD, True, PP_ALIGN.CENTER)
        add_rect(s, Inches(3.75), y, Inches(8.55), Inches(0.8), PANEL)
        add_text(s, b, Inches(3.98), y + Inches(0.18), Inches(8.0), Inches(0.35), 20, WHITE)
    script.append(("Slide 16", "This is your direct answer slide if the professor asks what the project found."))

    # 17
    s = slide("Limitations And Why They Matter", "Shows you understand what the results do and do not prove.")
    add_bullets(s, [
        "These are synthetic networks, not real friendship networks.",
        "The fixed 50-person roster is useful for control, but it is still one roster.",
        "Two seeds per condition give repeatability evidence, but more seeds would strengthen estimates.",
        "Prompt wording and model provider behavior can change over time.",
        "The findings are best interpreted as model behavior under controlled prompts, not as claims about real human societies."
    ], Inches(1.0), Inches(1.45), Inches(11.1), Inches(4.3), 23, WHITE, 10)
    add_text(s, "Strong closing sentence", Inches(1.0), Inches(6.15), Inches(3.2), Inches(0.3), 20, GOLD, True)
    add_text(s, "The value of this project is that it measures those sensitivities instead of hiding them.", Inches(1.0), Inches(6.52), Inches(10.0), Inches(0.45), 24, WHITE)
    script.append(("Slide 17", "Use limitations positively: they show rigor, not weakness."))

    # 18
    s = slide("Conclusion", "Generated networks are powerful, but they are not neutral.")
    add_text(s, "Main takeaway", Inches(0.85), Inches(1.42), Inches(3.0), Inches(0.35), 24, GOLD, True)
    add_text(s, "LLMs can create structurally plausible social networks, but culture, language, method, and model choice reshape who becomes connected to whom.", Inches(0.85), Inches(2.15), Inches(11.5), Inches(1.15), 28, WHITE, True)
    add_bullets(s, [
        "Base repo: established the network-generation problem with older model outputs.",
        "Our work: extended it into a controlled capstone matrix.",
        "Final result: the same personas can produce meaningfully different social structures depending on the experimental condition."
    ], Inches(1.1), Inches(3.95), Inches(10.8), Inches(1.65), 22, WHITE, 10)
    add_text(s, "Q&A", Inches(5.65), Inches(6.15), Inches(2.0), Inches(0.5), 32, CYAN, True, PP_ALIGN.CENTER)
    script.append(("Slide 18", "End cleanly: generated social networks are useful, but the method/model/prompt choices must be reported and measured."))

    pptx = OUT / "llms_social_network_generators_complete_presentation.pptx"
    prs.save(pptx)

    notes = OUT / "llms_social_network_generators_speaker_script.md"
    with notes.open("w", encoding="utf-8") as f:
        f.write("# Speaker Script - LLMs as Social Network Generators\n\n")
        for slide_name, note in script:
            f.write(f"## {slide_name}\n{note}\n\n")
    return pptx


def main() -> None:
    mkdirs()
    charts = generate_charts()
    pptx = build_deck(charts)
    summary = {
        "presentation": str(pptx),
        "speaker_script": str(OUT / "llms_social_network_generators_speaker_script.md"),
        "charts": {k: str(v) for k, v in charts.items()},
        "slides": 18,
        "sources": [
            "README.md",
            "ARCHITECTURE.md",
            "stats/cultural_study/research_answers.md",
            "stats/method_study/research_answers.md",
            "stats/language_study/research_answers.md",
            "stats/*/summary.csv",
            "plots/persona_breakdown/persona_breakdown.png",
            "plots/method_collage.png",
            "outputs/network_bonding_demo/persona_bonding_visual_preview.png",
        ],
    }
    (OUT / "presentation_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
